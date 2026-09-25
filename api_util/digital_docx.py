"""
api_util/digital_docx.py — the DOCX adapter of `digital_to_json.py` (Layer A, Issue #18).

python-docx (MIT) opens the OPC package and resolves table grids (merged cells); the body is
walked directly as WordprocessingML with lxml, because python-docx's `Paragraph.text` is
what lost text on the old path. What this module fixes, against the #10 §9 measurement
(gap register G3–G5) and the #18 mapping table:

* **Pages (G3).** A DOCX used to be one page. Breaks now follow atrium-alto-postprocess
  `text_formats._docx_handlers` / `_Flow` (#31) exactly, so the two tools number a
  document's pages the same way: `w:br w:type="page"`, `pageBreakBefore`, a section break
  that is not `continuous`/`nextColumn`, and — in the default `auto` mode — Word's own
  `w:lastRenderedPageBreak`, de-duplicated against an explicit break just before it. DOCX
  pages are still not rendered pages: without `lastRenderedPageBreak` (a file never saved
  by Word) only the explicit breaks exist.
* **Tracked changes (G5).** Text inside `w:ins` / `w:moveTo` is part of the document as it
  reads; `w:del` / `w:moveFrom` / `w:delText` are not. `Paragraph.text` dropped insertions.
* **Strict packages (G4).** python-docx refuses a package whose main part has no Override
  content type (alto-postprocess's `CTX000000010.docx`) — and, for the same check, every
  `.docm`, `.dotx` and `.dotm`. The package relationship still names the main part, so the
  content type is repaired in memory and the package reopened. `source.sha256` stays the
  hash of the file as given.
* **Structure.** Every paragraph is one `group_id` (`p{n}`), split into lines at `w:br` /
  `w:cr`; headings come from the outline level (the paragraph's, then its style chain's),
  then from the style name (`Heading N`, `Nadpis N`, `Title`, `Subtitle`); bold/italic are
  the effective values through run → character style → paragraph style; text boxes follow
  their anchor paragraph (`mc:Fallback` copies are skipped, so their text appears once);
  field codes (`instrText`) are skipped and field results kept.
* **Furniture and notes** (`lines[].style.region`). Each section's header and footer, once
  per section — the header on its first page, the footer on its last — and footnotes at the
  end of the page that references them, endnotes at the end of the document.

No geometry: a DOCX has none until something renders it, so no `bbox` and no `canvas`.
"""

from __future__ import annotations

import io
import os
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any, Callable, Dict, List, Optional, Tuple

from api_util.digital_ir import (
    REGION_FOOTER,
    REGION_FOOTNOTE,
    REGION_HEADER,
    TEXT_LAYER_BLANK,
    DigitalDocument,
    DigitalInputError,
    DigitalLine,
    DigitalPage,
    DigitalTable,
    missing_dependency,
    renumber_lines,
    sha256_file,
)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
WML_MAIN = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

#: How page breaks are read — alto-postprocess `[TEXT_INGEST].PAGE_BREAKS`, same values.
PAGE_BREAK_MODES = ("auto", "explicit", "none")

#: Zip limits, checked on DECLARED sizes before anything is decompressed (a zip bomb
#: declares its size honestly or fails the CRC). Generous for real reports.
ZIP_MAX_MEMBERS = 10_000
ZIP_MAX_TOTAL = 512 * 1024 * 1024
ZIP_MAX_RATIO = 200

#: Elements whose subtree is not document text. `pPr`/`rPr` are read by the handlers, not
#: walked; `Fallback` duplicates its `Choice`; the rest are deletions, field codes,
#: annotations and the marks inside a note that point back at its reference.
_SKIP = frozenset(
    {
        "del", "moveFrom", "delText", "delInstrText", "instrText", "rPr", "pPr", "sdtPr",
        "sdtEndPr", "fldData", "commentRangeStart", "commentRangeEnd", "commentReference",
        "bookmarkStart", "bookmarkEnd", "Fallback", "tblPr", "tblGrid", "trPr", "tcPr",
        "footnoteRef", "endnoteRef", "separator", "continuationSeparator", "annotationRef",
        "proofErr", "permStart", "permEnd", "sectPr",
    }
)  # fmt: skip

_HEADING_NAME = re.compile(
    r"^(?:heading|nadpis|überschrift|titre|t[ií]tulo|kop|nagłówek)\s*([1-9])$", re.IGNORECASE
)
_TITLE_NAMES = frozenset({"title", "název", "nazev", "titel", "titre", "tytuł"})
_SUBTITLE_NAMES = frozenset({"subtitle", "podtitul", "untertitel", "sous-titre", "podtytuł"})


def _local(tag: Any) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _attr(el: Any, name: str) -> Optional[str]:
    """A `w:`-namespaced attribute (or an un-namespaced one), else None."""
    value = el.get(f"{{{W_NS}}}{name}")
    return value if value is not None else el.get(name)


def _on(el: Any) -> bool:
    """ST_OnOff: an absent `w:val` means on."""
    value = _attr(el, "val")
    return value is None or value.lower() not in ("0", "false", "off", "none")


def _outline(el: Any) -> Optional[int]:
    """A `w:outlineLvl` value (0–9), or None when absent or malformed."""
    if el is None:
        return None
    try:
        return int(_attr(el, "val") or 9)
    except ValueError:
        return None


def _child(el: Any, name: str) -> Any:
    return next((c for c in el if _local(c.tag) == name), None) if el is not None else None


def _parse_xml(blob: bytes) -> Any:
    """Parse a package part with entities, DTDs and the network off."""
    from lxml import etree  # noqa: PLC0415  (python-docx dependency)

    parser = etree.XMLParser(
        resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False
    )
    return etree.fromstring(blob, parser)


# ── opening the package ───────────────────────────────────────────────────────


def _check_zip(path: str) -> None:
    try:
        with zipfile.ZipFile(path) as zf:
            infos = zf.infolist()
    except (zipfile.BadZipFile, OSError) as exc:
        raise DigitalInputError("corrupt", f"not a readable ZIP package ({exc})") from exc
    if len(infos) > ZIP_MAX_MEMBERS:
        raise DigitalInputError("zip_limits_exceeded", f"{len(infos)} members > {ZIP_MAX_MEMBERS}")
    total = sum(i.file_size for i in infos)
    packed = sum(i.compress_size for i in infos) or 1
    if total > ZIP_MAX_TOTAL or total / packed > ZIP_MAX_RATIO:
        raise DigitalInputError(
            "zip_limits_exceeded",
            f"declared {total} bytes uncompressed ({total / packed:.0f}x) — over the limits",
        )


def _main_part_name(zf: zipfile.ZipFile) -> str:
    """The main document part, from the package relationship — not from a guess."""
    try:
        rels = _parse_xml(zf.read("_rels/.rels"))
    except KeyError as exc:
        raise DigitalInputError("corrupt", "package has no _rels/.rels") from exc
    for rel in rels:
        if str(rel.get("Type", "")).endswith("/officeDocument"):
            return str(rel.get("Target", "")).lstrip("/")
    raise DigitalInputError("corrupt", "package relationships name no main document part")


def repaired_package(path: str) -> bytes:
    """The package with its main part's content type set to WordprocessingML main.

    Covers both a missing Override (the part falls back to `application/xml`) and the
    macro-enabled / template main types python-docx will not open. Every other entry is
    copied byte for byte; the repaired bytes live only in memory.
    """
    with zipfile.ZipFile(path) as zf:
        main = "/" + _main_part_name(zf)
        types = _parse_xml(zf.read("[Content_Types].xml"))
        for override in [o for o in types if _local(o.tag) == "Override"]:
            if override.get("PartName", "").lower() == main.lower():
                types.remove(override)
        from lxml import etree  # noqa: PLC0415

        etree.SubElement(types, f"{{{CT_NS}}}Override", PartName=main, ContentType=WML_MAIN)
        patched = etree.tostring(types, xml_declaration=True, encoding="UTF-8", standalone=True)
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as dst:
            for info in zf.infolist():
                data = patched if info.filename == "[Content_Types].xml" else zf.read(info)
                dst.writestr(info, data)
    return out.getvalue()


def open_package(path: str) -> Tuple[Any, bool]:
    """(python-docx Document, repaired?) — strict first, the in-memory repair second."""
    try:
        import docx  # noqa: PLC0415  (optional dependency, imported on use)
    except ImportError as exc:
        raise missing_dependency("python-docx", "docx") from exc

    _check_zip(path)
    try:
        return docx.Document(path), False
    except ValueError as exc:
        if "not a Word file" not in str(exc):
            raise DigitalInputError("corrupt", f"DOCX could not be opened ({exc})") from exc
    except KeyError as exc:
        raise DigitalInputError("corrupt", f"DOCX is missing a part ({exc})") from exc
    except Exception as exc:  # PackageNotFoundError, XMLSyntaxError, BadZipFile …
        raise DigitalInputError(
            "corrupt", f"DOCX could not be opened ({type(exc).__name__}: {exc})"
        ) from exc
    try:
        return docx.Document(io.BytesIO(repaired_package(path))), True
    except DigitalInputError:
        raise
    except Exception as exc:
        raise DigitalInputError(
            "corrupt", f"DOCX could not be opened ({type(exc).__name__}: {exc})"
        ) from exc


# ── styles ────────────────────────────────────────────────────────────────────


@dataclass
class _Style:
    name: str = ""
    based_on: Optional[str] = None
    outline: Optional[int] = None
    bold: Optional[bool] = None
    italic: Optional[bool] = None


@dataclass
class _Styles:
    """The styles part, read directly: python-docx's style API needs the part to carry its
    own content type, which a repaired package cannot promise."""

    by_id: Dict[str, _Style] = field(default_factory=dict)
    default_paragraph: Optional[str] = None

    @classmethod
    def load(cls, related: Dict[str, bytes]) -> "_Styles":
        out = cls()
        blob = related.get("styles")
        if not blob:
            return out
        root = _parse_xml(blob)
        for el in root:
            if _local(el.tag) != "style":
                continue
            style_id = _attr(el, "styleId") or ""
            name_el, based_el = _child(el, "name"), _child(el, "basedOn")
            ppr, rpr = _child(el, "pPr"), _child(el, "rPr")
            outline_el = _child(ppr, "outlineLvl")
            bold_el, italic_el = _child(rpr, "b"), _child(rpr, "i")
            out.by_id[style_id] = _Style(
                name=(_attr(name_el, "val") or "") if name_el is not None else "",
                based_on=_attr(based_el, "val") if based_el is not None else None,
                outline=_outline(outline_el),
                bold=_on(bold_el) if bold_el is not None else None,
                italic=_on(italic_el) if italic_el is not None else None,
            )
            if _attr(el, "type") == "paragraph" and _attr(el, "default") in ("1", "true"):
                out.default_paragraph = style_id
        return out

    def chain(self, style_id: Optional[str]) -> List[_Style]:
        seen, out = set(), []
        while style_id and style_id in self.by_id and style_id not in seen:
            seen.add(style_id)
            out.append(self.by_id[style_id])
            style_id = self.by_id[style_id].based_on
        return out

    def inherited(self, style_id: Optional[str], attr: str) -> Optional[Any]:
        for style in self.chain(style_id):
            value = getattr(style, attr)
            if value is not None:
                return value
        return None

    def heading_level(
        self, style_id: Optional[str], direct_outline: Optional[int]
    ) -> Optional[int]:
        """Outline level first (the paragraph's, then the style chain's), then the name."""
        outline = (
            direct_outline if direct_outline is not None else self.inherited(style_id, "outline")
        )
        if outline is not None:
            return min(outline + 1, 6) if 0 <= outline < 9 else None
        for style in self.chain(style_id)[:1]:
            for candidate in (style.name, style_id or ""):
                lowered = candidate.strip().lower()
                match = _HEADING_NAME.match(lowered) or re.match(
                    r"^(?:heading|nadpis)([1-9])$", lowered
                )
                if match:
                    return min(int(match.group(1)), 6)
                if lowered in _TITLE_NAMES:
                    return 1
                if lowered in _SUBTITLE_NAMES:
                    return 2
        return None


# ── the flow: text into lines, lines into pages ──────────────────────────────


@dataclass
class _Frag:
    text: str
    bold: bool
    italic: bool


class _Flow:
    """Lines into pages, in reading order — alto-postprocess `_Flow`'s semantics.

    `explicit_break()` starts a new page unless nothing was emitted yet (never an empty first
    page); `rendered_break()` (Word's lastRenderedPageBreak) only starts one if the current
    page already has text, which is what de-duplicates it against an explicit break just
    before it.
    """

    def __init__(self) -> None:
        self.pages: List[List[DigitalLine]] = [[]]
        self.buf: List[_Frag] = []
        self._any = False
        self.group: Optional[str] = None
        self.heading: Optional[int] = None
        self.region: Optional[str] = None

    def text(self, value: str, bold: bool, italic: bool) -> None:
        if value:
            self.buf.append(_Frag(value, bold, italic))

    def line_break(self) -> None:
        text = re.sub(r"\s+", " ", "".join(f.text for f in self.buf)).strip()
        inked = [f for f in self.buf if f.text.strip()]
        self.buf = []
        if not text:
            return
        self.pages[-1].append(
            DigitalLine(
                page="",
                line=0,
                text=text,
                bold=bool(inked) and all(f.bold for f in inked),
                italic=bool(inked) and all(f.italic for f in inked),
                heading_level=self.heading,
                group_id=self.group,
                region=self.region,
            )
        )
        self._any = True

    def explicit_break(self) -> None:
        self.line_break()
        if self._any:
            self.pages.append([])

    def rendered_break(self) -> None:
        self.line_break()
        if self.pages[-1]:
            self.pages.append([])

    @property
    def page_index(self) -> int:
        return len(self.pages) - 1


class _Walker:
    """Walks WordprocessingML into a `_Flow`. One instance per story (body, header, note)."""

    def __init__(
        self,
        flow: _Flow,
        styles: _Styles,
        mode: str,
        make_table: Callable[[Any], Any],
        group_prefix: str = "p",
        on_note: Optional[Callable[[str, str], None]] = None,
    ) -> None:
        self.flow = flow
        self.styles = styles
        self.mode = mode
        self.make_table = make_table
        self.group_prefix = group_prefix
        self.on_note = on_note
        self.paragraphs = 0
        self.tables: List[Tuple[int, DigitalTable]] = []
        self.table_count = 0
        self.cell_group: Optional[str] = None
        self.depth = 0
        self.deferred: List[Any] = []
        self.sections: List[Tuple[Any, int, int]] = []  # (sectPr, first page, last page)
        self._section_start = 0
        self._run: List[Tuple[bool, bool]] = []
        self._para_style: Optional[str] = None

    # -- helpers --------------------------------------------------------------

    def page_break(self) -> None:
        if self.mode == "none":
            self.flow.line_break()
        else:
            self.flow.explicit_break()

    def _run_flags(self, run: Any) -> Tuple[bool, bool]:
        rpr = _child(run, "rPr")
        style_id = None
        flags: List[Optional[bool]] = []
        for attr in ("b", "i"):
            direct = _child(rpr, attr)
            flags.append(_on(direct) if direct is not None else None)
        rstyle = _child(rpr, "rStyle")
        if rstyle is not None:
            style_id = _attr(rstyle, "val")
        bold = flags[0]
        italic = flags[1]
        if bold is None:
            bold = self.styles.inherited(style_id, "bold") if style_id else None
        if italic is None:
            italic = self.styles.inherited(style_id, "italic") if style_id else None
        if bold is None:
            bold = self.styles.inherited(self._para_style, "bold")
        if italic is None:
            italic = self.styles.inherited(self._para_style, "italic")
        return bool(bold), bool(italic)

    # -- the walk ---------------------------------------------------------------

    def walk_children(self, el: Any) -> None:
        for child in el:
            self.visit(child)

    def visit(self, el: Any) -> None:
        name = _local(el.tag)
        if not name or name in _SKIP:
            return
        # Handlers fire for WordprocessingML only: DrawingML reuses the local names `t`,
        # `br`, `p` and `r` (`a:t`, `a:br` …) for chart and SmartArt text, which is not
        # the document's own text flow.
        in_wml = isinstance(el.tag, str) and el.tag.startswith(f"{{{W_NS}}}")
        handler = getattr(self, f"_on_{name}", None) if in_wml else None
        if handler is not None:
            handler(el)
        else:
            self.walk_children(el)

    def _on_txbxContent(self, el: Any) -> None:
        self.deferred.append(el)

    def _on_p(self, el: Any) -> None:
        ppr = _child(el, "pPr")
        pstyle = _child(ppr, "pStyle")
        style_id = _attr(pstyle, "val") if pstyle is not None else self.styles.default_paragraph
        outline_el = _child(ppr, "outlineLvl")
        direct_outline = _outline(outline_el)
        break_before = _child(ppr, "pageBreakBefore")

        self.depth += 1
        self.flow.line_break()
        if break_before is not None and _on(break_before):
            self.page_break()
        saved = (self.flow.group, self.flow.heading, self._para_style)
        if self.cell_group is None:
            self.paragraphs += 1
            self.flow.group = f"{self.group_prefix}{self.paragraphs - 1}"
        else:
            self.flow.group = self.cell_group
        self.flow.heading = self.styles.heading_level(style_id, direct_outline)
        self._para_style = style_id
        self.walk_children(el)
        self.flow.line_break()
        self.flow.group, self.flow.heading, self._para_style = saved
        self.depth -= 1

        sect = _child(ppr, "sectPr")
        if sect is not None:
            self._end_section(sect)
        if self.depth == 0 and self.deferred:
            boxes, self.deferred = self.deferred, []
            for box in boxes:
                self.walk_children(box)

    def _end_section(self, sect: Any) -> None:
        self.sections.append((sect, self._section_start, self.flow.page_index))
        kind = _child(sect, "type")
        value = (_attr(kind, "val") or "") if kind is not None else ""
        if value not in ("continuous", "nextColumn"):
            self.page_break()
        self._section_start = self.flow.page_index

    def close_body(self, body: Any) -> None:
        """The body's own sectPr describes the last section."""
        sect = _child(body, "sectPr")
        self.flow.line_break()
        if sect is not None:
            self.sections.append((sect, self._section_start, self.flow.page_index))

    def _on_r(self, el: Any) -> None:
        self._run.append(self._run_flags(el))
        self.walk_children(el)
        self._run.pop()

    def _flags(self) -> Tuple[bool, bool]:
        return self._run[-1] if self._run else (False, False)

    def _on_t(self, el: Any) -> None:
        if el.text:
            self.flow.text(el.text, *self._flags())

    def _on_tab(self, el: Any) -> None:
        # Inside `w:tabs` (a paragraph property) a `w:tab` is a tab STOP, never text — but
        # pPr is skipped, so every `w:tab` reaching here is a run's tab character.
        self.flow.text(" ", *self._flags())

    def _on_ptab(self, el: Any) -> None:
        self.flow.text(" ", *self._flags())

    def _on_noBreakHyphen(self, el: Any) -> None:
        self.flow.text("-", *self._flags())

    def _on_softHyphen(self, el: Any) -> None:
        return

    def _on_br(self, el: Any) -> None:
        if (_attr(el, "type") or "") == "page":
            self.page_break()
        else:
            self.flow.line_break()

    def _on_cr(self, el: Any) -> None:
        self.flow.line_break()

    def _on_lastRenderedPageBreak(self, el: Any) -> None:
        if self.mode == "auto":
            self.flow.rendered_break()

    def _on_footnoteReference(self, el: Any) -> None:
        if self.on_note is not None:
            self.on_note("footnote", _attr(el, "id") or "")

    def _on_endnoteReference(self, el: Any) -> None:
        if self.on_note is not None:
            self.on_note("endnote", _attr(el, "id") or "")

    def _on_tbl(self, el: Any) -> None:
        """A table: grid shape to `tables[]`, cell text to `lines[]` grouped per cell.

        python-docx repeats a merged cell once per grid position it spans; each position
        still gets a `cells[]` entry, pointing at the first occurrence's group, and the text
        is walked once. Paragraphs inside a cell (nested tables included) carry the cell's
        group, which is the join key `tables[].cells[].group_id` names.
        """
        table = self.make_table(el)
        try:
            rows = list(table.rows)
        except Exception:
            rows = []
        try:
            n_cols = len(table.columns)
        except Exception:  # no w:tblGrid: python-docx cannot count columns, the rows can
            n_cols = 0
        if rows and not n_cols:
            try:
                n_cols = max(len(row.cells) for row in rows)
            except Exception:
                n_cols = 0
        if not rows or not n_cols:
            # `n_rows`/`n_cols` carry `minimum: 1`: a degenerate grid emitted verbatim would
            # make Layer D raise on our own output. It has no shape and no text to lose.
            return
        outer = self.cell_group is None
        number = self.table_count
        if outer:
            self.table_count += 1
        group = f"tbl{number}"
        grid = DigitalTable(
            table_id=f"t{number}", page="", n_rows=len(rows), n_cols=n_cols, group_id=group
        )
        start_page = self.flow.page_index
        origins: Dict[Any, str] = {}
        saved_cell = self.cell_group
        for row_no, row in enumerate(rows):
            for col_no, cell in enumerate(row.cells):
                first = origins.get(cell._tc)
                cell_group = first or f"{group}-r{row_no}c{col_no}"
                grid.cells.append(
                    {"row": row_no, "col": col_no, "is_header": row_no == 0, "group_id": cell_group}
                )
                if first is not None:
                    continue
                origins[cell._tc] = cell_group
                self.cell_group = cell_group if outer else saved_cell
                self.walk_children(cell._tc)
        self.cell_group = saved_cell
        if outer:
            self.tables.append((start_page, grid))


def _related_blobs(document: Any) -> Tuple[Dict[str, bytes], Dict[str, bytes]]:
    """({reltype-suffix: blob} for single parts, {rId: blob} for everything)."""
    by_type: Dict[str, bytes] = {}
    by_id: Dict[str, bytes] = {}
    for rel in document.part.rels.values():
        if rel.is_external:
            continue
        try:
            blob = rel.target_part.blob
        except Exception:
            continue
        by_id[rel.rId] = blob
        by_type.setdefault(rel.reltype.rsplit("/", 1)[-1], blob)
    return by_type, by_id


def _notes(blob: Optional[bytes]) -> Dict[str, Any]:
    """{w:id: note element}; separators are identified by `w:type`, not by id."""
    if not blob:
        return {}
    notes = {}
    for note in _parse_xml(blob):
        if _local(note.tag) not in ("footnote", "endnote"):
            continue
        if (_attr(note, "type") or "normal") != "normal":
            continue
        notes[_attr(note, "id") or ""] = note
    return notes


def _story_lines(
    element: Any, styles: _Styles, make_table: Callable[[Any], Any], prefix: str, region: str
) -> List[DigitalLine]:
    """Lines of a header, footer or note: its own flow, no page breaks, one region."""
    flow = _Flow()
    flow.region = region
    walker = _Walker(flow, styles, "none", make_table, group_prefix=prefix)
    walker.walk_children(element)
    flow.line_break()
    lines = [line for page in flow.pages for line in page]
    for line in lines:
        line.region = region
        line.heading_level = None
    return lines


def extract_docx(path: str, doc_id: str, origin: str, page_breaks: str = "auto") -> DigitalDocument:
    """Layer A for DOCX. Returns pages in reading order, lines not yet normalised."""
    if page_breaks not in PAGE_BREAK_MODES:
        raise ValueError(f"page_breaks must be one of {PAGE_BREAK_MODES}, got {page_breaks!r}")
    source, _repaired = open_package(path)

    from docx.table import Table  # noqa: PLC0415

    document = DigitalDocument(
        doc_id=doc_id,
        origin=origin,
        media_type=MEDIA_TYPE,
        sha256=sha256_file(path),
        filename=os.path.basename(path),
        reading_order="flow",
    )
    document.use("python-docx", "lxml")

    by_type, by_id = _related_blobs(source)
    styles = _Styles.load(by_type)

    def make_table(el: Any) -> Any:
        return Table(el, source)

    references: List[Tuple[str, str, int]] = []
    flow = _Flow()
    walker = _Walker(
        flow,
        styles,
        page_breaks,
        make_table,
        on_note=lambda kind, note_id: references.append((kind, note_id, flow.page_index)),
    )
    body = source.element.body
    walker.walk_children(body)
    walker.close_body(body)

    pages = flow.pages
    while len(pages) > 1 and not pages[-1]:
        pages.pop()
    n_pages = len(pages)
    ends: List[List[DigitalLine]] = [[] for _ in range(n_pages)]
    heads: List[List[DigitalLine]] = [[] for _ in range(n_pages)]
    feet: List[List[DigitalLine]] = [[] for _ in range(n_pages)]

    # Notes, at the end of the page that references them (endnotes: the last page).
    footnotes, endnotes = _notes(by_type.get("footnotes")), _notes(by_type.get("endnotes"))
    for kind, note_id, page_no in references:
        store, prefix = (footnotes, "fn") if kind == "footnote" else (endnotes, "en")
        note = store.pop(note_id, None)
        if note is None:
            continue
        target = min(page_no, n_pages - 1) if kind == "footnote" else n_pages - 1
        lines = _story_lines(note, styles, make_table, f"{prefix}{note_id}-", REGION_FOOTNOTE)
        for line in lines:
            line.group_id = f"{prefix}{note_id}"
        ends[target].extend(lines)

    # Headers and footers, once per section: the header on the section's first page, the
    # footer on its last. A section with no reference of its own is linked to the previous
    # one, and repeating the same text would only duplicate it.
    for number, (sect, first, last) in enumerate(walker.sections):
        first, last = min(first, n_pages - 1), min(last, n_pages - 1)
        title_page = _child(sect, "titlePg")
        use_first = title_page is not None and _on(title_page)
        for kind, region, store, page_no in (
            ("header", REGION_HEADER, heads, first),
            ("footer", REGION_FOOTER, feet, last),
        ):
            refs = {
                (_attr(r, "type") or "default"): r.get(f"{{{R_NS}}}id")
                for r in sect
                if _local(r.tag) == f"{kind}Reference"
            }
            lines: List[DigitalLine] = []
            for ref_type in (["first"] if use_first else []) + ["default"]:
                blob = by_id.get(refs.get(ref_type) or "")
                if blob:
                    prefix = "hdr" if kind == "header" else "ftr"
                    lines = _story_lines(
                        _parse_xml(blob), styles, make_table, f"{prefix}{number}-", region
                    )
                    if lines:
                        break
            store[page_no].extend(lines)

    for index in range(n_pages):
        label = str(index + 1)
        page = DigitalPage(page=label, page_index=index + 1, unit="pt")
        page.lines = heads[index] + pages[index] + ends[index] + feet[index]
        if not page.lines:
            page.text_layer = TEXT_LAYER_BLANK
        document.pages.append(page)
    for start_page, grid in walker.tables:
        page = document.pages[min(start_page, n_pages - 1)]
        grid.page = page.page
        page.tables.append(grid)

    renumber_lines(document)
    return document
